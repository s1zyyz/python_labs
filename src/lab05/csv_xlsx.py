import csv
from pathlib import Path
from openpyxl import Workbook
from openpyxl.utils import get_column_letter


def csv_to_xlsx(csv_path: str, xlsx_path: str) -> None:
    """
    Конвертирует CSV в XLSX с помощью openpyxl.
    Первая строка CSV — заголовок.
    Лист называется "Sheet1".
    Колонки — автоширина по длине текста (не менее 8 символов).
    """
    src = Path(csv_path)
    dst = Path(xlsx_path)

    
    if src.suffix.lower() != ".csv":
        raise ValueError(f"Ожидался входной файл .csv, получен: {csv_path}")
    if dst.suffix.lower() != ".xlsx":
        raise ValueError(f"Ожидался выходной файл .xlsx, получен: {xlsx_path}")

    
    if not src.exists():
        raise FileNotFoundError(f"Файл не найден: {csv_path}")

    
    if src.read_text(encoding="utf-8").strip() == "":
        raise ValueError("CSV-файл пуст")

    
    with src.open("r", encoding="utf-8", newline="") as f:
        reader = csv.reader(f)
        rows = list(reader)

    if not rows or not any(field.strip() for field in rows[0]):
        raise ValueError("CSV-файл пуст или не содержит заголовка")

    
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    for row in rows:
        ws.append(row)

    
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            cell_value = str(cell.value or "")
            if len(cell_value) > max_len:
                max_len = len(cell_value)
        ws.column_dimensions[col_letter].width = max(max_len + 2, 8)

    
    dst.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dst)


if __name__ == "__main__":
    
    csv_to_xlsx("data/people.csv", "data/out/people.xlsx")
    csv_to_xlsx("data/cities.csv", "data/out/cities.xlsx")
